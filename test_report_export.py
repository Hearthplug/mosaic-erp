"""Report export bytes: PDF parses with pypdf, XLSX with openpyxl, CSV with csv."""
import unittest

import report_export as rx


COLS = [('Account', 0.5, 'left'), ('Debit', 0.25, 'right'), ('Credit', 0.25, 'right')]
ROWS = [('1000 Cash', '$253.40', '$10.99'), ('1010 Bank', '$96.20', '$0.00')]
TOTALS = ('Totals', '$349.60', '$10.99')


class ReportExportBytes(unittest.TestCase):
    def test_pdf_round_trip(self):
        from pypdf import PdfReader
        import io
        pdf = rx.render_pdf('Trial balance', 'Test Shop - as of 2026-09-25', COLS, ROWS, TOTALS)
        self.assertTrue(pdf.startswith(b'%PDF-'))
        r = PdfReader(io.BytesIO(pdf))
        self.assertEqual(len(r.pages), 1)
        text = r.pages[0].extract_text()
        self.assertIn('Trial balance', text)
        self.assertIn('Test Shop', text)
        self.assertIn('$349.60', text)
        self.assertIn('Page 1 of 1', text)

    def test_pdf_paginates(self):
        from pypdf import PdfReader
        import io
        many = [('%04d Row' % i, '$1.00', '$2.00') for i in range(120)]
        pdf = rx.render_pdf('Long', 'Shop', COLS, many, TOTALS)
        r = PdfReader(io.BytesIO(pdf))
        self.assertGreaterEqual(len(r.pages), 3)
        self.assertIn('Page 2 of', r.pages[1].extract_text())

    def test_xlsx_round_trip(self):
        from openpyxl import load_workbook
        import io
        data = rx.render_xlsx('Trial balance', COLS, ROWS, TOTALS)
        ws = load_workbook(io.BytesIO(data)).active
        self.assertEqual([c.value for c in ws[1]], ['Account', 'Debit', 'Credit'])
        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws.cell(row=4, column=2).value, '$349.60')

    def test_csv_round_trip(self):
        import csv, io
        rows = list(csv.reader(io.StringIO(rx.render_csv(COLS, ROWS, TOTALS).decode())))
        self.assertEqual(rows[0], ['Account', 'Debit', 'Credit'])
        self.assertEqual(len(rows), 4)

    def test_pdf_escapes_parens(self):
        pdf = rx.render_pdf('A (test) \\ report', 'S', COLS, ROWS)
        self.assertTrue(pdf.startswith(b'%PDF-'))


if __name__ == '__main__':
    unittest.main()
