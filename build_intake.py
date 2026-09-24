"""File intake for the Build screen: PDFs, spreadsheets, and text.

Deterministic extraction only - no model calls in this module. Photos go
through the owner's own vision key (ChatProviderClient.extract_image).
Everything returned here is an excerpt plus a suggestion; nothing is staged,
drafted, or applied until the owner clicks, and every draft still passes the
normal review screen before activation.
"""
import base64, csv, io

from migration_packs import SCHEMAS

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARS = 200_000
EXCERPT_CHARS = 4_000
SHEET_ROW_CAP = 5_000

TEXT_TYPES = {'text/plain', 'text/csv', 'text/markdown', 'application/csv'}
SHEET_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PDF_TYPE = 'application/pdf'
PHOTO_TYPES = {'image/png', 'image/jpeg', 'image/webp'}


def decode_upload(data_b64):
    try:
        raw = base64.b64decode(data_b64 or '', validate=True)
    except Exception:
        raise ValueError('The file could not be read. Try saving it again and re-uploading it.')
    if not raw:
        raise ValueError('That file is empty.')
    if len(raw) > MAX_BYTES:
        raise ValueError('Files up to 5 MB are supported. Split the file and try again.')
    return raw


def sniff_mime(name, declared):
    low = (name or '').lower()
    if low.endswith('.pdf'):
        return PDF_TYPE
    if low.endswith('.xlsx'):
        return SHEET_TYPE
    if low.endswith('.csv'):
        return 'text/csv'
    if low.endswith(('.txt', '.md')):
        return 'text/plain'
    for ext, mime in (('.png', 'image/png'), ('.jpg', 'image/jpeg'), ('.jpeg', 'image/jpeg'), ('.webp', 'image/webp')):
        if low.endswith(ext):
            return mime
    return declared or ''


def _sheet_to_csv(raw):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError('Spreadsheet reading is not installed on this server.')
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ValueError('That spreadsheet could not be opened. Save it as .xlsx or .csv and try again.')
    if not wb.worksheets:
        raise ValueError('That spreadsheet has no sheets.')
    out = io.StringIO()
    writer = csv.writer(out)
    rows = 0
    for row in wb.worksheets[0].iter_rows(values_only=True):
        writer.writerow(['' if cell is None else str(cell) for cell in row])
        rows += 1
        if rows >= SHEET_ROW_CAP:
            break
    return out.getvalue()


def _pdf_text(raw):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ValueError('PDF reading is not installed on this server.')
    try:
        reader = PdfReader(io.BytesIO(raw))
    except Exception:
        raise ValueError('That PDF could not be opened. If it is a photo scan, photograph it instead.')
    parts = []
    for page in reader.pages[:8]:
        try:
            parts.append(page.extract_text() or '')
        except Exception:
            pass
    text = '\n'.join(p for p in parts if p).strip()
    if not text:
        raise ValueError('No readable text in that PDF - it is probably a scan. Photograph it with the phone camera instead.')
    return text


def _detect_sheet_pack(text):
    """Score the header row against the migration pack schemas."""
    first = text.splitlines()[0] if text else ''
    try:
        header = next(csv.reader(io.StringIO(first)))
    except Exception:
        return None, []
    cols = {c.strip().lower() for c in header if c and c.strip()}
    if not cols:
        return None, []
    best, best_hits, missing = None, 0, []
    for pack, (required, _key) in SCHEMAS.items():
        hits = len(required & cols)
        if hits > best_hits:
            best, best_hits = pack, hits
            missing = sorted(required - cols)
    if best and best_hits >= max(2, len(SCHEMAS[best][0]) - 1):
        return best, missing
    return None, []


def _document_hints(text):
    low = text.lower()
    hints = []
    if 'invoice' in low or 'bill to' in low:
        hints.append('invoice')
    if 'report' in low or 'summary' in low:
        hints.append('report')
    if 'dashboard' in low:
        hints.append('dashboard')
    return hints


def read_file(name, declared_mime, data_b64):
    """Extract an excerpt and a suggestion from an uploaded document.
    Raises ValueError with an owner-facing message on failure."""
    raw = decode_upload(data_b64)
    mime = sniff_mime(name, declared_mime)
    if mime == PDF_TYPE:
        text, kind = _pdf_text(raw), 'pdf'
    elif mime == SHEET_TYPE:
        text, kind = _sheet_to_csv(raw), 'sheet'
    elif mime in TEXT_TYPES or not mime:
        text = raw.decode('utf-8', errors='replace')
        kind = 'sheet' if mime == 'text/csv' else 'text'
    else:
        raise ValueError('That file type is not supported. Use a PDF, an Excel sheet, a CSV, or a text file.')
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
    out = {'name': name, 'kind': kind, 'characters': len(text), 'excerpt': text[:EXCERPT_CHARS]}
    if kind == 'sheet':
        pack, missing = _detect_sheet_pack(text)
        if pack:
            out['detected'] = {'type': 'migration', 'pack': pack, 'missing_columns': missing, 'csv': text}
        else:
            out['detected'] = {'type': 'unknown', 'csv': text}
    else:
        out['detected'] = {'type': 'document', 'hints': _document_hints(text)}
    return out
