"""Report exports: CSV, XLSX and PDF for every report in the app.

PDF is generated with a small built-in writer (no external library):
core Helvetica font, A4 pages, header with the shop name and date
range, page numbers, and simple column layout with right-aligned
numbers. XLSX uses openpyxl. Every export returns real bytes the
owner can open.
"""

import csv
import io
import time

PAGE_W, PAGE_H = 595, 842  # A4 portrait in points
MARGIN = 48
ROW_H = 16


def _esc(text):
    return str(text).replace('\\', r'\\').replace('(', r'\(').replace(')', r'\)')


def _pdf_stream(lines):
    return ('\n'.join(lines)).encode('latin-1', errors='replace')


def render_pdf(title, subtitle, columns, rows, totals=None, money_cols=None):
    """Build a real PDF document. columns: list of (label, width_share, align).
    rows: list of string lists. totals: optional final row. Returns bytes."""
    money_cols = money_cols or set()
    # paginate: header block + rows per page
    usable_h = PAGE_H - 2 * MARGIN - 70  # header (title+subtitle+rule+col heads)
    per_page = int(usable_h // ROW_H)
    pages = []
    body = list(rows) + ([totals] if totals else [])
    for i in range(0, max(len(body), 1), per_page):
        pages.append(body[i:i + per_page])
    objects = []  # (body bytes); object ids 1..N
    # 1 catalog, 2 pages, 3 font; then per page: page obj + content obj
    page_ids, content_ids = [], []
    next_id = 4
    for _ in pages:
        page_ids.append(next_id); content_ids.append(next_id + 1); next_id += 2
    objects.append(b'<< /Type /Catalog /Pages 2 0 R >>')
    kids = ' '.join('%d 0 R' % p for p in page_ids)
    objects.append(('<< /Type /Pages /Kids [%s] /Count %d >>' % (kids, len(pages))).encode())
    objects.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
    total = len(pages)
    for idx, chunk in enumerate(pages):
        objects.append(('<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] '
                        '/Resources << /Font << /F1 3 0 R /F2 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >> >> >> '
                        '/Contents %d 0 R >>' % (PAGE_W, PAGE_H, content_ids[idx])).encode())
        lines = []
        y = PAGE_H - MARGIN
        # header
        lines.append('BT /F2 14 Tf %d %d Td (%s) Tj ET' % (MARGIN, y, _esc(title)))
        if subtitle:
            y -= 18
            lines.append('BT /F1 9 Tf %d %d Td (%s) Tj ET' % (MARGIN, y, _esc(subtitle)))
        y -= 8
        lines.append('%d %d m %d %d l S' % (MARGIN, y, PAGE_W - MARGIN, y))
        # column headers
        y -= ROW_H
        x = MARGIN
        col_w = [(PAGE_W - 2 * MARGIN) * share for _l, share, _a in columns]
        for ci, (label, share, align) in enumerate(columns):
            if align == 'right':
                est = len(label) * 6.2
                lines.append('BT /F2 9 Tf %d %d Td (%s) Tj ET' % (int(max(x, x + col_w[ci] - est)), y, _esc(label)))
            else:
                lines.append('BT /F2 9 Tf %d %d Td (%s) Tj ET' % (int(x), y, _esc(label)))
            x += col_w[ci]
        lines.append('%d %d m %d %d l S' % (MARGIN, y - 4, PAGE_W - MARGIN, y - 4))
        y -= ROW_H
        for row in chunk:
            x = MARGIN
            is_total = totals is not None and row is totals
            font = '/F2' if is_total else '/F1'
            for ci, cell in enumerate(row):
                cell = str(cell)
                if columns[ci][2] == 'right':
                    est = len(cell) * 5.6
                    lines.append('BT %s 9 Tf %d %d Td (%s) Tj ET' % (font, int(max(x, x + col_w[ci] - est)), y, _esc(cell)))
                else:
                    lines.append('BT %s 9 Tf %d %d Td (%s) Tj ET' % (font, int(x), y, _esc(cell)))
                x += col_w[ci]
            if is_total:
                lines.append('%d %d m %d %d l S' % (MARGIN, y - 4, PAGE_W - MARGIN, y - 4))
            y -= ROW_H
        lines.append('BT /F1 8 Tf %d %d Td (Page %d of %d) Tj ET' % (PAGE_W - MARGIN - 60, MARGIN - 16, idx + 1, total))
        stream = _pdf_stream(lines)
        objects.append(b'<< /Length %d >>\nstream\n' % len(stream) + stream + b'\nendstream')
    # assemble with xref
    out = io.BytesIO()
    out.write(b'%PDF-1.4\n')
    offsets = []
    for i, body_b in enumerate(objects):
        offsets.append(out.tell())
        out.write(('%d 0 obj\n' % (i + 1)).encode() + body_b + b'\nendobj\n')
    xref_at = out.tell()
    out.write(('xref\n0 %d\n' % (len(objects) + 1)).encode())
    out.write(b'0000000000 65535 f \n')
    for off in offsets:
        out.write(('%010d 00000 n \n' % off).encode())
    out.write(('trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF' % (len(objects) + 1, xref_at)).encode())
    return out.getvalue()


def render_csv(columns, rows, totals=None):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[0] for c in columns])
    for r in rows:
        w.writerow(list(r))
    if totals:
        w.writerow(list(totals))
    return buf.getvalue().encode('utf-8')


def render_xlsx(title, columns, rows, totals=None):
    from copy import copy
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = title[:28] or 'Report'
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        f=copy(cell.font); f.bold=True; cell.font=f
    for r in rows:
        ws.append(list(r))
    if totals:
        ws.append(list(totals))
        for cell in ws[ws.max_row]:
            f=copy(cell.font); f.bold=True; cell.font=f
    for ci, col in enumerate(columns, start=1):
        width = max(len(str(col[0])), max((len(str(r[ci - 1])) for r in rows), default=8)) + 2
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = min(width, 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SYMBOLS = {'USD': '$', 'EUR': '€', 'GBP': '£', 'INR': '₹'}


def _money(minor, currency):
    sym = SYMBOLS.get(currency, currency + ' ')
    return '%s%.2f' % (sym, (minor or 0) / 100.0)


def _shop(store, wid):
    row = store._db.execute('SELECT name FROM workspaces WHERE id=?', (wid,)).fetchone()
    return (row['name'] if row else None) or 'Your shop'


def build_report(report, wid, qs, store, books, dayclose_summary):
    """Return (title, subtitle, columns, rows, totals) for a report id.
    columns: (label, width_share, align). Raises ValueError for an unknown report."""
    shop = _shop(store, wid)
    cur = books.status(wid).get('base_currency', 'USD')
    as_of = qs.get('as_of', [None])[0]
    date_from = qs.get('from', [None])[0]
    date_to = qs.get('to', [None])[0]
    span = ''
    if date_from or date_to:
        span = ' - %s to %s' % (date_from or 'start', date_to or 'today')
    elif as_of:
        span = ' - as of %s' % as_of

    if report == 'trial-balance':
        tb = books.trial_balance(wid, as_of)
        rows = [('%s %s' % (a['code'], a['name']), _money(a['debit_minor'], cur), _money(a['credit_minor'], cur)) for a in tb['accounts']]
        totals = ('Totals', _money(tb['total_debit_minor'], cur), _money(tb['total_credit_minor'], cur))
        return ('Trial balance', shop + span, [('Account', 0.5, 'left'), ('Debit', 0.25, 'right'), ('Credit', 0.25, 'right')], rows, totals)

    if report == 'profit-loss':
        st = books.financial_statements(wid, to_date=as_of)
        rows = [('%s %s' % (a['code'], a['name']), _money(a['amount_minor'], cur)) for a in st['profit_and_loss']]
        totals = ('Net profit', _money(st['net_profit_minor'], cur))
        return ('Profit and loss', shop + span, [('Account', 0.7, 'left'), ('Amount', 0.3, 'right')], rows, totals)

    if report == 'balance-sheet':
        st = books.financial_statements(wid, to_date=as_of)
        rows = [('%s %s' % (a['code'], a['name']), a['type'].capitalize(), _money(a['amount_minor'], cur)) for a in st['balance_sheet']]
        return ('Balance sheet', shop + span, [('Account', 0.5, 'left'), ('Type', 0.2, 'left'), ('Amount', 0.3, 'right')], rows, None)

    if report in ('bills', 'invoices'):
        kind = 'purchase_bill' if report == 'bills' else 'sales_invoice'
        docs = books.list_documents(wid, kind, '500').get('documents', [])
        docs = [d for d in docs if d.get('kind') == kind]
        rows = [(d['number'], d.get('issue_date') or '', d.get('party_name') or '', d.get('status', ''), _money(d['total_minor'], d.get('currency') or cur), _money(d['balance_minor'], d.get('currency') or cur)) for d in docs]
        tot_t = sum(d['total_minor'] for d in docs)
        tot_b = sum(d['balance_minor'] for d in docs)
        title = 'Supplier bills' if report == 'bills' else 'Customer invoices'
        return (title, shop + span, [('Number', 0.15, 'left'), ('Date', 0.15, 'left'), ('Party', 0.25, 'left'), ('Status', 0.12, 'left'), ('Total', 0.16, 'right'), ('Balance', 0.17, 'right')], rows, ('Totals', '', '', '', _money(tot_t, cur), _money(tot_b, cur)))

    if report in ('aging-payable', 'aging-receivable'):
        kind = 'payable' if report == 'aging-payable' else 'receivable'
        import datetime
        ag = books.aging(wid, as_of or datetime.date.today().isoformat(), kind)
        rows = [(d['number'], d.get('due_date') or '', str(d.get('days_overdue', 0)), d.get('bucket', ''), _money(d['balance_minor'], d.get('currency') or cur)) for d in ag.get('documents', [])]
        b = ag.get('buckets', {})
        totals = ('Buckets: current %s' % _money(b.get('current', 0), cur), '1-30: %s' % _money(b.get('1_30', 0), cur), '31-60: %s' % _money(b.get('31_60', 0), cur), '61+: %s' % _money(b.get('61_90', 0) + b.get('over_90', 0), cur), '')
        title = 'What you owe suppliers' if kind == 'payable' else 'What customers owe you'
        return (title + ' (aging)', shop + span, [('Number', 0.2, 'left'), ('Due date', 0.2, 'left'), ('Days overdue', 0.2, 'right'), ('Bucket', 0.15, 'left'), ('Balance', 0.25, 'right')], rows, totals)

    if report == 'stock-register':
        rows_raw = store._db.execute("SELECT p.sku,p.name,p.unit,l.name AS location_name,COALESCE((SELECT SUM(sl.quantity_delta) FROM stock_ledger sl WHERE sl.workspace_id=p.workspace_id AND sl.product_id=p.id AND sl.location_id=l.id),0) AS on_hand FROM retail_products p JOIN locations l ON l.workspace_id=p.workspace_id AND l.active=1 WHERE p.workspace_id=? AND p.active=1 ORDER BY p.sku,l.code", (wid,)).fetchall()
        rows = [(r['sku'], r['name'], r['location_name'], str(r['on_hand']), r['unit']) for r in rows_raw]
        return ('Stock register', shop + span, [('SKU', 0.15, 'left'), ('Item', 0.35, 'left'), ('Store', 0.25, 'left'), ('On hand', 0.15, 'right'), ('Unit', 0.1, 'left')], rows, None)

    if report == 'sales':
        sql = "SELECT s.number,s.sold_at,s.status,s.currency,s.total_minor,l.code AS location_code FROM sales s JOIN locations l ON l.id=s.location_id WHERE s.workspace_id=?"
        args = [wid]
        if date_from:
            sql += " AND date(s.sold_at)>=date(?)"; args.append(date_from)
        if date_to:
            sql += " AND date(s.sold_at)<=date(?)"; args.append(date_to)
        sql += " ORDER BY s.sold_at DESC LIMIT 500"
        rows_raw = store._db.execute(sql, args).fetchall()
        rows = [(r['number'], (r['sold_at'] or '').replace('T',' ')[:16], r['location_code'], r['status'], _money(r['total_minor'], r['currency'] or cur)) for r in rows_raw]
        tot = sum(r['total_minor'] for r in rows_raw)
        return ('Sales', shop + span, [('Bill', 0.18, 'left'), ('When', 0.27, 'left'), ('Store', 0.15, 'left'), ('Status', 0.15, 'left'), ('Total', 0.25, 'right')], rows, ('Totals', '', '', '', _money(tot, cur)))

    if report == 'tills':
        rows_raw = store._db.execute("SELECT cs.opened_at,cs.closed_at,cs.opening_minor,cs.expected_minor,cs.actual_minor,cs.variance_minor,cs.status,l.code AS location_code FROM cash_sessions cs JOIN locations l ON l.id=cs.location_id WHERE cs.workspace_id=? ORDER BY cs.opened_at DESC LIMIT 200", (wid,)).fetchall()
        rows = [(r['location_code'], (r['opened_at'] or '').replace('T',' ')[:16], (r['closed_at'] or '-').replace('T',' ')[:16], r['status'], _money(r['opening_minor'], cur), _money(r['expected_minor'], cur), _money(r['actual_minor'], cur) if r['actual_minor'] is not None else '-', _money(r['variance_minor'], cur) if r['variance_minor'] is not None else '-') for r in rows_raw]
        return ('Till sessions', shop + span, [('Store', 0.09, 'left'), ('Opened', 0.18, 'left'), ('Closed', 0.18, 'left'), ('Status', 0.1, 'left'), ('Float', 0.12, 'right'), ('Expected', 0.12, 'right'), ('Counted', 0.11, 'right'), ('Over/short', 0.1, 'right')], rows, None)

    if report == 'day-close':
        day = qs.get('date', [None])[0]
        dc = dayclose_summary(wid, day)
        rows = [('Sales total', _money(dc['sales_total_minor'], cur)), ('Bills sold', str(dc['sales_count'])), ('Items sold', str(dc['items_sold'])), ('Cash taken', _money(dc['payments_cash_minor'], cur)), ('Card/bank taken', _money(dc['payments_bank_minor'], cur)), ('Sold on credit', _money(dc['credit_minor'], cur)), ('Cash expected in drawers', _money(dc['expected_cash_minor'], cur))]
        return ('Day close', '%s - %s' % (shop, dc.get('date', day or 'today')), [('What', 0.6, 'left'), ('Amount', 0.4, 'right')], rows, None)

    raise ValueError('Unknown report: %s' % report)


REPORTS = ['trial-balance', 'profit-loss', 'balance-sheet', 'bills', 'invoices', 'aging-payable', 'aging-receivable', 'stock-register', 'sales', 'tills', 'day-close']
